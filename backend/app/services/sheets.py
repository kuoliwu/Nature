"""通用 Excel / CSV 匯出匯入框架。

每種資料（建案、檢查清單、時段、預約名單、人員、驗屋結果…）只要定義：
  - 欄位（表頭、型別、是否必填、是否只匯出）
  - 匯出時的資料列
  - 匯入時如何寫回（apply）
共同規則：以「關鍵欄位」比對更新、新的新增、檔案裡沒有的不刪除；先檢查全部，
有任何錯誤整份不寫入（同一個交易內 rollback），dry_run 只回報不寫入。
"""

import csv
import io
import re
from collections.abc import Callable
from dataclasses import dataclass, field
from datetime import date, datetime, time
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from pydantic import EmailStr, TypeAdapter, ValidationError

from app.services.units_excel import ExcelFormatError, _csv_rows, _xlsx_rows

_email = TypeAdapter(EmailStr)
TRUE_WORDS = {"是", "y", "yes", "true", "1", "v", "✓", "開放", "啟用", "o"}
FALSE_WORDS = {"否", "n", "no", "false", "0", "x", "關閉", "停用"}
MAX_ROWS = 5000


@dataclass
class Col:
    key: str
    header: str
    width: int = 14
    kind: str = "text"  # text / int / date / time / bool / email / phone / choice
    required: bool = False
    export_only: bool = False  # 匯出時提供參考，匯入時忽略
    choices: dict[str, Any] | None = None  # 顯示文字 → 值（kind=choice）
    max_len: int = 255

    @property
    def title(self) -> str:
        return self.header + ("*" if self.required else "")


@dataclass
class ImportOutcome:
    created: list[str] = field(default_factory=list)
    updated: list[str] = field(default_factory=list)
    unchanged: int = 0
    errors: list[tuple[int, str]] = field(default_factory=list)

    def error(self, row: int, message: str) -> None:
        self.errors.append((row, message))


@dataclass
class Dataset:
    name: str
    sheet: str  # 工作表名稱，也用於檔名
    columns: list[Col]
    instructions: list[str]
    export_rows: Callable[..., list[dict]]  # (db, ctx) -> rows
    apply: Callable[..., None]  # (db, ctx, rows: list[(row_no, dict)], outcome) -> None
    scope: str = "global"  # global / project / project_service / batch
    export_permission: str | None = None  # None：登入即可
    import_permission: str = "project.manage"
    text_keys: tuple[str, ...] = ()  # 匯出時設成文字格式的欄位（保留開頭的 0）


# ---------------------------------------------------------------------------
# 匯出
# ---------------------------------------------------------------------------
def _display(col: Col, value: Any) -> Any:
    if value is None:
        return None
    if col.kind == "bool":
        return "是" if value else "否"
    if col.kind == "choice" and col.choices:
        reverse = {v: k for k, v in col.choices.items()}
        return reverse.get(value, value)
    if col.kind == "time" and isinstance(value, time):
        return value.strftime("%H:%M")
    return value


def export_xlsx(ds: Dataset, rows: list[dict]) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = ds.sheet[:31]
    fill = PatternFill("solid", fgColor="DBEAFE")
    ref_fill = PatternFill("solid", fgColor="F3F4F6")
    for c, col in enumerate(ds.columns, start=1):
        cell = ws.cell(row=1, column=c, value=col.title)
        cell.font = Font(bold=True, color="6B7280" if col.export_only else "000000")
        cell.fill = ref_fill if col.export_only else fill
        cell.alignment = Alignment(horizontal="center")
        ws.column_dimensions[get_column_letter(c)].width = col.width
    ws.freeze_panes = "A2"
    text_cols = [i for i, col in enumerate(ds.columns, start=1) if col.kind in ("phone", "text", "time") or col.key in ds.text_keys]
    for r, row in enumerate(rows, start=2):
        for c, col in enumerate(ds.columns, start=1):
            cell = ws.cell(row=r, column=c, value=_display(col, row.get(col.key)))
            if col.kind == "date" and isinstance(cell.value, date):
                cell.number_format = "yyyy-mm-dd"
            elif c in text_cols:
                cell.number_format = "@"
    for r in range(len(rows) + 2, len(rows) + 202):  # 空白列預設文字格式，新增時電話不會被吃掉 0
        for c in text_cols:
            ws.cell(row=r, column=c).number_format = "@"

    help_ws = wb.create_sheet("說明")
    help_ws.column_dimensions["A"].width = 100
    lines = ["填寫說明", "", *ds.instructions, "",
             "・表頭有 * 的欄位為必填；灰色表頭的欄位只供參考，匯入時會忽略。",
             "・匯入前系統會先檢查每一列，有任何錯誤時整份都不會寫入，並列出錯誤的列號。",
             "・檔案裡沒有的資料不會被刪除。",
             "・也可以存成 CSV（UTF-8 或 Big5）匯入，第一列同樣是表頭。"]
    for i, line in enumerate(lines, start=1):
        help_ws.cell(row=i, column=1, value=line).font = Font(bold=(i == 1), size=14 if i == 1 else 11)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def export_csv(ds: Dataset, rows: list[dict]) -> bytes:
    buf = io.StringIO()
    writer = csv.writer(buf, lineterminator="\r\n")
    writer.writerow([c.title for c in ds.columns])
    for row in rows:
        out = []
        for col in ds.columns:
            v = _display(col, row.get(col.key))
            out.append("" if v is None else v.isoformat() if isinstance(v, date) else v)
        writer.writerow(out)
    return ("﻿" + buf.getvalue()).encode("utf-8")


# ---------------------------------------------------------------------------
# 匯入：解析
# ---------------------------------------------------------------------------
def _text(v) -> str | None:
    if v is None:
        return None
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    s = str(v).strip()
    return s or None


def convert(col: Col, raw: Any) -> Any:
    """把儲存格內容轉成欄位型別；格式錯誤丟 ValueError（訊息給使用者看）"""
    text = _text(raw)
    if text is None:
        return None
    name = col.header
    if col.kind == "int":
        try:
            return int(float(text))
        except ValueError:
            raise ValueError(f"{name}「{text}」不是整數")
    if col.kind == "date":
        if isinstance(raw, datetime):
            return raw.date()
        if isinstance(raw, date):
            return raw
        for fmt in ("%Y-%m-%d", "%Y/%m/%d", "%Y.%m.%d"):
            try:
                return datetime.strptime(text, fmt).date()
            except ValueError:
                pass
        raise ValueError(f"{name}「{text}」格式不正確，請用 2026-12-31")
    if col.kind == "time":
        if isinstance(raw, time):
            return raw.replace(second=0, microsecond=0)
        if isinstance(raw, datetime):
            return raw.time().replace(second=0, microsecond=0)
        if isinstance(raw, float) and 0 <= raw < 1:  # Excel 時間是一天的比例
            minutes = round(raw * 24 * 60)
            return time(minutes // 60, minutes % 60)
        m = re.fullmatch(r"(\d{1,2})[:：](\d{2})(?::\d{2})?", text)
        if not m or int(m.group(1)) > 23 or int(m.group(2)) > 59:
            raise ValueError(f"{name}「{text}」格式不正確，請用 09:30")
        return time(int(m.group(1)), int(m.group(2)))
    if col.kind == "bool":
        low = text.lower()
        if low in TRUE_WORDS:
            return True
        if low in FALSE_WORDS:
            return False
        raise ValueError(f"{name}「{text}」請填「是」或「否」")
    if col.kind == "email":
        try:
            _email.validate_python(text)
        except ValidationError:
            raise ValueError(f"{name}「{text}」格式不正確")
        return text
    if col.kind == "phone":
        if isinstance(raw, (int, float)) and re.fullmatch(r"9\d{8}", text):
            text = "0" + text  # Excel 吃掉的開頭 0
        return text[: col.max_len]
    if col.kind == "choice" and col.choices is not None:
        if text not in col.choices:
            raise ValueError(f"{name}「{text}」請填：{'、'.join(col.choices)}")
        return col.choices[text]
    if len(text) > col.max_len:
        raise ValueError(f"{name}超過 {col.max_len} 個字")
    return text


def parse(ds: Dataset, data: bytes) -> tuple[list[tuple[int, dict]], list[str], list[tuple[int, str]]]:
    """回傳 (資料列[(列號, 值)], 檔案中出現的欄位 key, 錯誤[(列號, 訊息)])"""
    if data.startswith(b"\xd0\xcf\x11\xe0"):
        raise ExcelFormatError("這是舊版 .xls，請在 Excel 另存為 .xlsx 或 CSV 後再匯入")
    rows_iter = _xlsx_rows(data) if data.startswith(b"PK") else _csv_rows(data)
    header = next(rows_iter, None) or ()
    by_header = {c.header.replace(" ", ""): c for c in ds.columns if not c.export_only}
    index: dict[int, Col] = {}
    for i, title in enumerate(header):
        col = by_header.get(str(title or "").strip().rstrip("*").replace(" ", ""))
        if col:
            index[i] = col
    present = {c.key for c in index.values()}
    missing = [c.header for c in ds.columns if c.required and not c.export_only and c.key not in present]
    if missing:
        raise ExcelFormatError(f"找不到必要欄位：{'、'.join(missing)}。請使用匯出的範本，第一列表頭不要修改")

    rows, errors = [], []
    for row_no, values in enumerate(rows_iter, start=2):
        if row_no - 1 > MAX_ROWS:
            raise ExcelFormatError(f"一次最多匯入 {MAX_ROWS} 列")
        raw = {col.key: (values[i] if i < len(values) else None) for i, col in index.items()}
        if all(_text(v) is None for v in raw.values()):
            continue
        item, row_errors = {}, []
        for key, value in raw.items():
            col = next(c for c in index.values() if c.key == key)
            try:
                item[key] = convert(col, value)
            except ValueError as e:
                row_errors.append(str(e))
        for col in index.values():
            if col.required and item.get(col.key) is None and not any(col.header in e for e in row_errors):
                row_errors.append(f"{col.header}不能空白")
        if row_errors:
            errors.append((row_no, "；".join(row_errors)))
        else:
            rows.append((row_no, item))
    return rows, sorted(present), errors
