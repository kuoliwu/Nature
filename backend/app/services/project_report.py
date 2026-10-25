"""建案報表匯出（主管會議用）：互動網頁 HTML（單一檔案）與 Excel（含原生圖表）"""

import io
import json
from pathlib import Path

from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

TEMPLATE = Path(__file__).resolve().parent.parent / "templates" / "project_report.html"

# 與網頁版同一組色票（dataviz 驗證過）
STAGE_COLORS = {"completed": "1C5CAB", "booked": "3987E5", "invited": "86B6EF", "not_invited": "CDE2FB"}
SERVICE_COLORS = {"visit": "2A78D6", "inspection": "EB6834", "health_check": "1BAF7A"}
RESULT_COLORS = {"pass": "0CA30C", "fail": "D03B3B", "na": "C3C2B7"}
SERVICES = [("visit", "參觀"), ("inspection", "驗屋"), ("health_check", "健檢")]
HEADER_FILL = PatternFill("solid", fgColor="DBEAFE")


def build_html(data: dict) -> bytes:
    # json.dumps 後把 < 跳脫，避免資料內容跳出 <script>
    payload = json.dumps(data, ensure_ascii=False).replace("<", "\\u003c")
    return TEMPLATE.read_text("utf-8").replace("__DATA__", payload).encode("utf-8")


def _table(ws, top: int, headers: list[str], rows: list[list]) -> tuple[int, int]:
    """從第 top 列寫表格；回傳 (第一筆資料列, 最後一筆資料列)"""
    for c, h in enumerate(headers, start=1):
        cell = ws.cell(row=top, column=c, value=h)
        cell.font = Font(bold=True)
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(horizontal="center")
    for r, row in enumerate(rows, start=top + 1):
        for c, v in enumerate(row, start=1):
            ws.cell(row=r, column=c, value=v)
    return top + 1, top + len(rows)


def _color_series(chart: BarChart, colors: list[str]) -> None:
    for series, color in zip(chart.series, colors):
        series.graphicalProperties.solidFill = color
        series.graphicalProperties.line.solidFill = color


def _bar_chart(title: str, horizontal: bool, stacked: bool) -> BarChart:
    chart = BarChart()
    chart.type = "bar" if horizontal else "col"
    chart.title = title
    chart.style = 2
    if stacked:
        chart.grouping = "stacked"
        chart.overlap = 100
    chart.gapWidth = 60
    chart.height, chart.width = 8, 18
    chart.legend.position = "b"
    return chart


def build_xlsx(data: dict) -> bytes:
    wb = Workbook()
    units_total = data["units"]["total"]

    # --- 總覽：重點數字 + 各階段進度 ---
    ws = wb.active
    ws.title = "總覽"
    ws["A1"] = f"{data['project']['name']}　建案報表"
    ws["A1"].font = Font(bold=True, size=16)
    ws["A2"] = f"{data['project']['address']}　資料時間 {data['generated_at'].replace('T', ' ')}"
    ws["A2"].font = Font(color="6B7280")
    kpis = [["住戶數", units_total], ["缺 Email", data["units"]["without_email"]], ["逾期未完成報告", data["overdue_reports"]]]
    for p in data["phases"]:
        kpis.append([f"{p['label']}完成", p["stages"]["completed"]])
        if p["has_reports"]:
            kpis.append([f"{p['label']}報告已寄", p["reports_sent"]])
    _table(ws, 4, ["項目", "數量"], kpis)

    top = 6 + len(kpis)
    ws.cell(row=top - 1, column=1, value="各階段進度（戶）").font = Font(bold=True, size=13)
    stage_keys = list(STAGE_COLORS)
    first, last = _table(
        ws, top, ["階段", *[data["stage_labels"][k] for k in stage_keys], "完成率"],
        [[p["label"], *[p["stages"][k] for k in stage_keys],
          f"{round(p['stages']['completed'] / units_total * 100) if units_total else 0}%"] for p in data["phases"]],
    )
    chart = _bar_chart("各階段進度", horizontal=True, stacked=True)
    chart.add_data(Reference(ws, min_col=2, max_col=1 + len(stage_keys), min_row=top, max_row=last), titles_from_data=True)
    chart.set_categories(Reference(ws, min_col=1, min_row=first, max_row=last))
    chart.y_axis.title = "戶數"
    _color_series(chart, list(STAGE_COLORS.values()))
    ws.add_chart(chart, "H4")

    # --- 每日預約 ---
    ws2 = wb.create_sheet("每日預約")
    rows = [[d["date"], *[d[k] for k, _ in SERVICES], sum(d[k] for k, _ in SERVICES)] for d in data["daily"]]
    first, last = _table(ws2, 1, ["日期", *[label for _, label in SERVICES], "合計"], rows or [["（尚無預約）", 0, 0, 0, 0]])
    chart = _bar_chart("每日預約數", horizontal=False, stacked=True)
    chart.add_data(Reference(ws2, min_col=2, max_col=4, min_row=1, max_row=last), titles_from_data=True)
    chart.set_categories(Reference(ws2, min_col=1, min_row=first, max_row=last))
    chart.y_axis.title = "組數"
    _color_series(chart, list(SERVICE_COLORS.values()))
    ws2.add_chart(chart, "G2")

    # --- 檢查結果 ---
    ws3 = wb.create_sheet("檢查結果")
    result_rows = []
    for key, label in (("inspection", "驗屋"), ("health_check", "健檢")):
        v = data["results"][key]
        total = v["pass"] + v["fail"] + v["na"]
        result_rows.append([label, v["pass"], v["fail"], v["na"], f"{round(v['fail'] / total * 100) if total else 0}%"])
    first, last = _table(ws3, 1, ["", "合格", "不合格", "不適用", "不合格率"], result_rows)
    chart = _bar_chart("檢查結果", horizontal=True, stacked=True)
    chart.add_data(Reference(ws3, min_col=2, max_col=4, min_row=1, max_row=last), titles_from_data=True)
    chart.set_categories(Reference(ws3, min_col=1, min_row=first, max_row=last))
    _color_series(chart, list(RESULT_COLORS.values()))
    chart.height = 6
    ws3.add_chart(chart, "G2")

    # --- 不合格 Top 10 ---
    ws4 = wb.create_sheet("不合格Top10")
    fails = [[f"{r['service_label']}｜{r['label']}", r["count"]] for r in reversed(data["top_fails"])]  # 長條圖由下往上畫
    first, last = _table(ws4, 1, ["項目", "次數"], fails or [["（目前沒有不合格項目）", 0]])
    chart = _bar_chart("最常見的不合格項目", horizontal=True, stacked=False)
    chart.add_data(Reference(ws4, min_col=2, min_row=1, max_row=last), titles_from_data=True)
    chart.set_categories(Reference(ws4, min_col=1, min_row=first, max_row=last))
    chart.legend = None  # 單一系列不需要圖例
    _color_series(chart, [SERVICE_COLORS["visit"]])
    chart.height = max(6, 1 + 0.8 * len(fails))
    ws4.add_chart(chart, "D2")

    for sheet in wb.worksheets:
        for col in range(1, 8):
            sheet.column_dimensions[get_column_letter(col)].width = 14
        sheet.column_dimensions["A"].width = 34 if sheet.title == "不合格Top10" else 16
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
