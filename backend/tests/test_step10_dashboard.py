"""建案儀表板"""

from datetime import date, datetime, time, timedelta

import pytest

from app.database import SessionLocal
from app.models import Booking, BookingBatch, CheckResult, Inspection, InspectionResult, InspectionStatus, ServiceType, Unit
from app.services.slots import generate_slots
from tests.conftest import ADMIN


@pytest.fixture(autouse=True)
def _login(admin_token):
    pass


def test_dashboard_counts(client):
    pid = client.post("/projects", headers=ADMIN, json={"name": "儀表板建案", "address": "台南市"}).json()["id"]
    for i in range(5):
        client.post(f"/projects/{pid}/units", headers=ADMIN, json={"unit_no": f"D-{i}", "owner_email": f"d{i}@example.com" if i < 4 else None})

    # 參觀：寄邀請給 4 戶，2 戶預約，其中 1 戶已報到
    visit = client.post(f"/projects/{pid}/batches", headers=ADMIN, json={
        "name": "參觀", "service_type": "visit", "start_date": str(date.today() + timedelta(days=1)),
        "end_date": str(date.today() + timedelta(days=1)), "day_start": "10:00", "day_end": "11:00",
    }).json()
    client.post(f"/batches/{visit['id']}/invites/send", headers=ADMIN, json={})
    links = {r["unit_no"]: r["link"].rsplit("/", 1)[1] for r in client.get(f"/batches/{visit['id']}/invites", headers=ADMIN).json() if r["link"]}
    slot = visit["days"][0]["slots"][0]["id"]
    for n in ("D-0", "D-1"):
        client.post(f"/api/public/book/{links[n]}", json={"slot_id": slot, "contact_name": n, "contact_phone": "0900000000", "contact_email": f"{n}@example.com"})
    bid = next(r for r in client.get(f"/batches/{visit['id']}/invites", headers=ADMIN).json() if r["unit_no"] == "D-0")["booking"]["id"]
    client.post(f"/bookings/{bid}/checkin", headers=ADMIN)

    # 驗屋：1 戶已簽名，含 1 個不合格
    with SessionLocal() as db:
        unit = db.query(Unit).filter_by(project_id=pid, unit_no="D-2").one()
        batch = BookingBatch(project_id=pid, name="驗屋", service_type=ServiceType.INSPECTION)
        db.add(batch)
        db.flush()
        s = generate_slots(db, batch, date.today() + timedelta(days=2), time(9), time(9, 30))[0]
        b = Booking(batch_id=batch.id, slot_id=s.id, unit_id=unit.id, contact_name="x", contact_phone="0900")
        db.add(b)
        db.flush()
        db.add(Inspection(booking_id=b.id, status=InspectionStatus.SIGNED, signed_at=datetime.now()))
        point_id = client.get(f"/projects/{pid}/checklist", headers=ADMIN).json()[0]["items"][0]["check_points"][0]["id"]
        db.add(InspectionResult(booking_id=b.id, check_point_id=point_id, result=CheckResult.FAIL))
        db.commit()

    d = client.get(f"/projects/{pid}/dashboard", headers=ADMIN).json()
    assert d["units"] == {"total": 5, "with_email": 4, "without_email": 1}
    phases = {p["service"]: p for p in d["phases"]}
    assert phases["visit"]["stages"] == {"completed": 1, "booked": 1, "invited": 2, "not_invited": 1}
    assert phases["inspection"]["stages"]["completed"] == 1 and phases["inspection"]["reports_sent"] == 0
    assert phases["health_check"]["stages"] == {"completed": 0, "booked": 0, "invited": 0, "not_invited": 5}
    assert phases["visit"]["has_reports"] is False
    assert sum(day["visit"] for day in d["daily"]) == 2 and sum(day["inspection"] for day in d["daily"]) == 1
    assert d["results"]["inspection"]["fail"] == 1
    assert d["top_fails"][0]["label"].startswith("窗戶／") and d["top_fails"][0]["count"] == 1


def test_report_export_html_and_xlsx(client):
    import io

    from openpyxl import load_workbook

    pid = client.post("/projects", headers=ADMIN, json={"name": "報表</script><script>alert(1)</script>"}).json()["id"]
    client.post(f"/projects/{pid}/units", headers=ADMIN, json={"unit_no": "R-1"})

    link = client.post(f"/projects/{pid}/report-export-link", headers=ADMIN, params={"format": "html"}).json()
    assert link["filename"].startswith("建案報表_") and link["filename"].endswith(".html")
    html = client.get("/" + link["url"].split("://", 1)[1].split("/", 1)[1]).text  # 不需登入 header
    assert "const DATA = " in html and '"total": 1' in html
    assert "</script><script>alert(1)" not in html  # 建案名稱不能跳出 script

    link = client.post(f"/projects/{pid}/report-export-link", headers=ADMIN, params={"format": "xlsx"}).json()
    data = client.get("/" + link["url"].split("://", 1)[1].split("/", 1)[1]).content
    wb = load_workbook(io.BytesIO(data))
    assert wb.sheetnames == ["總覽", "每日預約", "檢查結果", "不合格Top10"]
    assert all(len(ws._charts) == 1 for ws in wb.worksheets)
    assert client.post(f"/projects/{pid}/report-export-link", headers=ADMIN, params={"format": "pdf"}).status_code == 422
