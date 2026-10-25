"""各種資料的 Excel / CSV 匯出匯入"""

import io
from datetime import date, timedelta

import pytest
from openpyxl import load_workbook

from tests.conftest import ADMIN, ADMIN_LOGIN

DAY = date.today() + timedelta(days=3)


@pytest.fixture(autouse=True)
def _login(admin_token):
    pass


def export(client, name: str, fmt: str = "xlsx", **params):
    r = client.post(f"/sheets/{name}/export-link", headers=ADMIN, params={"format": fmt, **params})
    assert r.status_code == 200, r.text
    data = client.get("/" + r.json()["url"].split("://", 1)[1].split("/", 1)[1]).content
    if fmt == "csv":
        return data
    return load_workbook(io.BytesIO(data))


def upload(client, name: str, wb_or_bytes, dry_run=False, headers=ADMIN, **params):
    if not isinstance(wb_or_bytes, bytes):
        buf = io.BytesIO()
        wb_or_bytes.save(buf)
        wb_or_bytes = buf.getvalue()
    return client.post(
        f"/sheets/{name}/import", headers=headers, params={"dry_run": dry_run, **params},
        files={"file": ("data.xlsx", wb_or_bytes)},
    )


def rows(wb) -> list[list]:
    return [list(r) for r in wb.active.iter_rows(values_only=True) if any(v is not None for v in r)]


def append(wb, values: dict):
    """依表頭名稱加一列"""
    ws = wb.active
    header = [str(c.value).rstrip("*") for c in ws[1]]
    row = last_row(wb) + 1  # 範本預留的空白列已設格式，ws.append 會加到更下面
    for c, h in enumerate(header, start=1):
        ws.cell(row=row, column=c, value=values.get(h))


def last_row(wb) -> int:
    ws = wb.active
    return max(r for r in range(1, ws.max_row + 1) if any(c.value is not None for c in ws[r]))


def set_cell(wb, row: int, header: str, value):
    ws = wb.active
    col = [str(c.value).rstrip("*") for c in ws[1]].index(header) + 1
    ws.cell(row=row, column=col).value = value  # ws.cell(value=None) 不會清空


def setup_batch(client, name: str, service="inspection", units=("A-1", "A-2")):
    pid = client.post("/projects", headers=ADMIN, json={"name": name}).json()["id"]
    for u in units:
        client.post(f"/projects/{pid}/units", headers=ADMIN, json={"unit_no": u, "owner_name": f"屋主{u}", "owner_phone": "0911222333"})
    batch = client.post(f"/projects/{pid}/batches", headers=ADMIN, json={
        "name": "第一梯", "service_type": service, "start_date": str(DAY), "end_date": str(DAY),
        "day_start": "09:00", "day_end": "10:00",
    }).json()
    return pid, batch["id"]


def test_projects_round_trip(client):
    client.post("/projects", headers=ADMIN, json={"name": "匯出建案", "phone": "062223333"})
    wb = export(client, "projects")
    assert wb.sheetnames == ["建案", "說明"]
    head = rows(wb)[0]
    assert head[0] == "案名*" and "住戶數" in head
    row_no = next(i for i, r in enumerate(rows(wb), start=1) if r[0] == "匯出建案")
    set_cell(wb, row_no, "社區經理", "王經理")
    append(wb, {"案名": "匯入新建案", "建案電話": "0912345678"})

    preview = upload(client, "projects", wb, dry_run=True).json()
    assert preview["dry_run"] and preview["created"] == ["匯入新建案"] and preview["updated"] == ["匯出建案"]
    assert not any(p["name"] == "匯入新建案" for p in client.get("/projects", headers=ADMIN).json())

    r = upload(client, "projects", wb).json()
    assert not r["dry_run"] and not r["errors"]
    new = next(p for p in client.get("/projects", headers=ADMIN).json() if p["name"] == "匯入新建案")
    assert new["phone"] == "0912345678"
    # 新建案自動帶入兩份檢查清單
    assert client.get(f"/projects/{new['id']}/checklist", headers=ADMIN, params={"service": "health_check"}).json()


def test_errors_mean_nothing_written(client):
    wb = export(client, "projects")
    append(wb, {"案名": "不該建立"})
    append(wb, {"建案電話": "123"})  # 缺案名
    r = upload(client, "projects", wb).json()
    assert r["dry_run"] and r["errors"][0]["message"] == "案名不能空白"
    assert not any(p["name"] == "不該建立" for p in client.get("/projects", headers=ADMIN).json())

    r = client.post("/sheets/projects/import", headers=ADMIN, files={"file": ("x.csv", "名稱,電話\r\nA,1".encode())})
    assert r.status_code == 422 and "案名" in r.json()["detail"]


def test_checklist_import_adds_and_disables(client):
    pid = client.post("/projects", headers=ADMIN, json={"name": "清單匯入"}).json()["id"]
    wb = export(client, "checklist", project_id=pid, service="health_check")
    first = rows(wb)[1]
    set_cell(wb, 2, "啟用", "否")
    append(wb, {"類別": "新類別", "項目": "新項目", "檢查點": "新檢查點"})
    r = upload(client, "checklist", wb, project_id=pid, service="health_check").json()
    assert r["created"] == ["新類別／新項目／新檢查點"] and len(r["updated"]) == 1, r

    cats = client.get(f"/projects/{pid}/checklist", headers=ADMIN, params={"service": "health_check"}).json()
    point = next(p for c in cats if c["name"] == first[0] for i in c["items"] if i["name"] == first[1]
                 for p in i["check_points"] if p["name"] == first[2])
    assert point["is_active"] is False
    assert any(c["name"] == "新類別" for c in cats)
    # 驗屋清單不受影響
    insp = client.get(f"/projects/{pid}/checklist", headers=ADMIN, params={"service": "inspection"}).json()
    assert not any(c["name"] == "新類別" for c in insp)


def test_slots_import(client):
    pid, bid = setup_batch(client, "時段匯入")
    wb = export(client, "slots", batch_id=bid)
    assert rows(wb)[0][:3] == ["日期*", "開始*", "結束"]
    set_cell(wb, 2, "名額", 5)
    set_cell(wb, 3, "開放", "否")
    append(wb, {"日期": DAY, "開始": "13:00", "結束": "13:30", "名額": 2})
    append(wb, {"日期": str(DAY), "開始": "14:00"})  # 新時段沒有結束時間
    r = upload(client, "slots", wb, batch_id=bid).json()
    assert r["dry_run"] and r["errors"] == [{"row": 5, "message": "新增時段需要填「結束」時間"}]

    wb.active.delete_rows(5)
    r = upload(client, "slots", wb, batch_id=bid).json()
    assert not r["errors"] and r["created"] == [f"{DAY} 13:00"] and len(r["updated"]) == 2
    slots = client.get(f"/batches/{bid}", headers=ADMIN).json()["days"][0]["slots"]
    assert slots[0]["capacity"] == 5 and slots[1]["is_open"] is False and slots[-1]["start_time"].startswith("13:00")


def test_bookings_import_respects_capacity(client):
    pid, bid = setup_batch(client, "預約匯入", units=("B-1", "B-2", "B-3", "B-4"))
    wb = export(client, "bookings", batch_id=bid)
    assert [r[0] for r in rows(wb)[1:]] == ["B-1", "B-2", "B-3", "B-4"] and rows(wb)[1][8] == "未預約"
    for i in range(2, 6):  # 4 戶都塞 09:00（名額 3）
        set_cell(wb, i, "日期", DAY)
        set_cell(wb, i, "開始時間", "09:00")
    r = upload(client, "bookings", wb, batch_id=bid).json()
    assert r["dry_run"] and "超過名額" in r["errors"][0]["message"]

    set_cell(wb, 5, "開始時間", "09:30")
    set_cell(wb, 5, "電話", 912000111)  # Excel 吃掉開頭 0
    r = upload(client, "bookings", wb, batch_id=bid).json()
    assert not r["errors"] and len(r["created"]) == 4
    invites = {i["unit_no"]: i for i in client.get(f"/batches/{bid}/invites", headers=ADMIN).json()}
    assert invites["B-1"]["booking"]["contact_name"] == "屋主B-1"  # 沒填聯絡人時用屋主
    assert invites["B-4"]["booking"]["contact_phone"] == "0912000111"

    # 改期：B-1 改到 09:30
    wb = export(client, "bookings", batch_id=bid)
    assert rows(wb)[1][8] == "已預約"
    set_cell(wb, 2, "開始時間", "09:30")
    r = upload(client, "bookings", wb, batch_id=bid).json()
    assert r["updated"] == ["B-1"] and r["unchanged"] == 3


def test_users_import(client):
    wb = export(client, "users")
    head = rows(wb)[0]
    assert "手機*" in head and "新密碼" in head and "建案 / 戶別管理" in head
    append(wb, {"姓名": "匯入人員", "手機": 912777888, "身分": "一般人員", "檢視預約時間表": "是"})
    r = upload(client, "users", wb).json()
    assert r["errors"][0]["message"] == "新帳號 0912777888 需要填「新密碼」"

    set_cell(wb, last_row(wb), "新密碼", "short")
    assert upload(client, "users", wb).json()["errors"][0]["message"] == "新密碼至少 8 碼"
    set_cell(wb, last_row(wb), "新密碼", "password-123")
    r = upload(client, "users", wb).json()
    assert r["created"] == ["匯入人員（0912777888）"], r
    login = client.post("/auth/login", json={"phone": "0912777888", "password": "password-123"}).json()
    assert login["user"]["permissions"] == ["booking.view"]

    # 不能停用自己
    wb = export(client, "users")
    me = next(i for i, r in enumerate(rows(wb), start=1) if r[1] == ADMIN_LOGIN["phone"])
    set_cell(wb, me, "啟用", "否")
    assert upload(client, "users", wb).json()["errors"][0]["message"] == "不能停用或降級自己的帳號"

    # 沒有人員管理權限不能匯入
    staff = {"Authorization": f"Bearer {login['access_token']}"}
    assert upload(client, "users", wb, headers=staff).status_code == 403


def test_results_import(client):
    pid, bid = setup_batch(client, "結果匯入", units=("C-1",))
    wb = export(client, "bookings", batch_id=bid)
    set_cell(wb, 2, "日期", DAY)
    set_cell(wb, 2, "開始時間", "09:00")
    assert not upload(client, "bookings", wb, batch_id=bid).json()["errors"]

    wb = export(client, "results", batch_id=bid)
    data = rows(wb)
    assert data[0][:5] == ["戶別*", "預約日期", "類別*", "項目*", "檢查點*"] and len(data) > 5
    set_cell(wb, 2, "結果", "不合格")
    set_cell(wb, 2, "備註", "牆面裂縫")
    set_cell(wb, 3, "結果", "合格")
    set_cell(wb, 4, "結果", "壞掉")
    r = upload(client, "results", wb, batch_id=bid).json()
    assert r["errors"][0]["row"] == 4 and "合格、不合格" in r["errors"][0]["message"]
    set_cell(wb, 4, "結果", None)
    r = upload(client, "results", wb, batch_id=bid).json()
    assert not r["errors"] and len(r["created"]) == 2

    booking_id = client.get(f"/batches/{bid}/invites", headers=ADMIN).json()[0]["booking"]["id"]
    insp = client.get(f"/inspections/{booking_id}", headers=ADMIN).json()
    pts = [p for c in insp["categories"] for i in c["items"] for p in i["points"]]
    assert pts[0]["result"] == "fail" and pts[0]["note"] == "牆面裂縫" and pts[1]["result"] == "pass"

    # 屋主簽名後鎖定
    from app.database import SessionLocal
    from app.models import Inspection, InspectionStatus

    with SessionLocal() as db:
        db.query(Inspection).filter_by(booking_id=booking_id).update({"status": InspectionStatus.SIGNED})
        db.commit()
    unchanged = export(client, "results", batch_id=bid)  # 已簽名但原封不動匯回不是錯誤
    assert upload(client, "results", unchanged, batch_id=bid).json()["errors"] == []
    set_cell(wb, 2, "結果", "合格")
    assert "已簽名" in upload(client, "results", wb, batch_id=bid).json()["errors"][0]["message"]

    # 每一種資料：匯出後原封不動匯回 → 沒有錯誤、沒有變動
    for name, params in [("projects", {}), ("users", {}), ("checklist", {"project_id": pid, "service": "inspection"}),
                         ("slots", {"batch_id": bid}), ("bookings", {"batch_id": bid}), ("results", {"batch_id": bid})]:
        for fmt in ("xlsx", "csv"):
            r = upload(client, name, export(client, name, fmt=fmt, **params), **params).json()
            assert (r["errors"], r["created"], r["updated"]) == ([], [], []), (name, fmt, r)


def test_csv_export_and_import(client):
    data = export(client, "projects", fmt="csv")
    assert data.startswith("﻿".encode()) and "案名*".encode() in data
    csv = "案名,建案電話\r\nCSV建案,0912000000\r\n".encode("big5")
    r = upload(client, "projects", csv).json()
    assert r["created"] == ["CSV建案"]


def test_visit_has_no_results_sheet(client):
    pid, bid = setup_batch(client, "參觀無結果", service="visit")
    r = client.post("/sheets/results/export-link", headers=ADMIN, params={"batch_id": bid})
    assert r.status_code == 422
