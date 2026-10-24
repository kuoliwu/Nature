"""住戶資料 Excel 匯出 / 匯入"""

import io
from datetime import datetime

import pytest
from openpyxl import Workbook, load_workbook

from tests.conftest import ADMIN


@pytest.fixture(autouse=True)
def _login(admin_token):
    pass


@pytest.fixture
def project_id(client) -> int:
    pid = client.post("/projects", headers=ADMIN, json={"name": "Excel 測試建案"}).json()["id"]
    client.post(f"/projects/{pid}/units", headers=ADMIN, json={
        "unit_no": "A-2F-1", "building": "A", "floor": "2F", "owner_name": "王小明",
        "owner_phone": "0912345678", "owner_email": "wang@example.com", "area_ping": 30.5, "handover_date": "2026-12-20",
    })
    client.post(f"/projects/{pid}/units", headers=ADMIN, json={"unit_no": "A-2F-2", "owner_name": "李小華"})
    return pid


def export(client, pid) -> Workbook:
    r = client.get(f"/projects/{pid}/units/export", headers=ADMIN)
    assert r.status_code == 200
    assert r.headers["content-type"].startswith("application/vnd.openxmlformats")
    return load_workbook(io.BytesIO(r.content))


def upload(client, pid, wb: Workbook | bytes, dry_run=False, headers=ADMIN):
    data = wb if isinstance(wb, bytes) else _to_bytes(wb)
    return client.post(
        f"/projects/{pid}/units/import", headers=headers, params={"dry_run": dry_run},
        files={"file": ("units.xlsx", data, "application/octet-stream")},
    )


def _to_bytes(wb: Workbook) -> bytes:
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def rows(ws) -> list[list]:
    return [list(r) for r in ws.iter_rows(min_row=2, values_only=True) if any(v is not None for v in r)]


def test_export_contains_units_and_instructions(client, project_id):
    wb = export(client, project_id)
    assert wb.sheetnames == ["住戶", "說明"]
    ws = wb["住戶"]
    assert [c.value for c in ws[1]][:3] == ["戶別*", "棟別", "樓層"]
    first = rows(ws)[0]
    assert first[0] == "A-2F-1" and first[7] == "0912345678" and first[8] == "wang@example.com"
    assert first[5].date().isoformat() == "2026-12-20"  # 交屋日期是 Excel 日期

    empty_pid = client.post("/projects", headers=ADMIN, json={"name": "空白建案"}).json()["id"]
    assert rows(export(client, empty_pid)["住戶"]) == []  # 沒有住戶時就是空白範本


def test_round_trip_update_and_create(client, project_id):
    wb = export(client, project_id)
    ws = wb["住戶"]
    ws["I3"] = "lee@example.com"  # 第 3 列（A-2F-2）補 Email
    ws.append(["B-3F-1", "B", "3F", "某路 3 號", 28, datetime(2027, 1, 5), "陳先生", 987654321, None, "新戶"])

    preview = upload(client, project_id, wb, dry_run=True).json()
    assert preview == {"dry_run": True, "created": ["B-3F-1"], "updated": ["A-2F-2"], "unchanged": 1, "errors": []}
    assert len(client.get(f"/projects/{project_id}/units", headers=ADMIN).json()) == 2  # 預覽不寫入

    result = upload(client, project_id, wb).json()
    assert result["dry_run"] is False and result["created"] == ["B-3F-1"]
    units = {u["unit_no"]: u for u in client.get(f"/projects/{project_id}/units", headers=ADMIN).json()}
    assert units["A-2F-2"]["owner_email"] == "lee@example.com"
    new = units["B-3F-1"]
    assert (new["owner_phone"], new["handover_date"], new["area_ping"], new["building"]) == ("0987654321", "2027-01-05", 28, "B")

    # 再匯入一次：全部都沒變
    again = upload(client, project_id, export(client, project_id)).json()
    assert (again["created"], again["updated"], again["unchanged"]) == ([], [], 3)


def test_blank_cell_clears_and_missing_columns_untouched(client, project_id):
    wb = Workbook()
    ws = wb.active
    ws.append(["戶別", "屋主Email"])  # 只有兩欄
    ws.append(["A-2F-1", None])
    upload(client, project_id, wb)
    unit = next(u for u in client.get(f"/projects/{project_id}/units", headers=ADMIN).json() if u["unit_no"] == "A-2F-1")
    assert unit["owner_email"] is None  # 有這欄、留白 → 清空
    assert unit["owner_name"] == "王小明"  # 沒有這欄 → 不動


def test_errors_reported_and_nothing_written(client, project_id):
    wb = Workbook()
    ws = wb.active
    ws.append(["戶別*", "屋主Email", "交屋日期", "坪數"])
    ws.append(["C-1", "ok@example.com", "2026/12/01", "30"])  # 正確
    ws.append(["C-2", "not-an-email", None, None])           # 第 3 列
    ws.append(["C-3", None, "明年", None])                   # 第 4 列
    ws.append([None, None, None, None])                      # 空白列略過
    ws.append(["C-1", None, None, None])                     # 第 6 列：重複
    ws.append([None, "x@example.com", None, "abc"])          # 第 7 列：沒有戶別

    r = upload(client, project_id, wb).json()
    assert r["dry_run"] is True  # 有錯誤 → 沒有寫入
    assert [e["row"] for e in r["errors"]] == [3, 4, 6, 7]
    messages = " ".join(e["message"] for e in r["errors"])
    assert "Email" in messages and "交屋日期" in messages and "重複" in messages and "戶別不能空白" in messages
    assert not any(u["unit_no"] == "C-1" for u in client.get(f"/projects/{project_id}/units", headers=ADMIN).json())


def test_bad_files(client, project_id):
    r = upload(client, project_id, b"this is not excel")  # 純文字會被當成 CSV 讀，但沒有戶別表頭
    assert r.status_code == 422 and "戶別" in r.json()["detail"]
    wb = Workbook()
    wb.active.append(["名稱", "電話"])
    r = upload(client, project_id, wb)
    assert r.status_code == 422 and "戶別" in r.json()["detail"]


def test_permissions(client, project_id):
    client.post("/users", headers=ADMIN, json={"name": "唯讀2", "phone": "0966000222", "password": "readonly-2"})
    token = client.post("/auth/login", json={"phone": "0966000222", "password": "readonly-2"}).json()["access_token"]
    h = {"Authorization": f"Bearer {token}"}
    assert client.get(f"/projects/{project_id}/units/export", headers=h).status_code == 200
    assert upload(client, project_id, export(client, project_id), headers=h).status_code == 403


def test_download_link_without_login(client, project_id):
    link = client.post(f"/projects/{project_id}/units/export-link", headers=ADMIN).json()
    assert link["filename"].startswith("住戶_Excel 測試建案_") and link["filename"].endswith(".xlsx")
    path = link["url"].split("://", 1)[1].split("/", 1)[1]
    r = client.get(f"/{path}")  # 不帶登入 header，靠連結裡的 token
    assert r.status_code == 200 and r.content.startswith(b"PK")

    csv_link = client.post(f"/projects/{project_id}/units/export-link", headers=ADMIN, params={"format": "csv"}).json()
    r = client.get("/" + csv_link["url"].split("://", 1)[1].split("/", 1)[1])
    text = r.content.decode("utf-8-sig")
    assert text.splitlines()[0].startswith("戶別*,棟別,樓層") and "A-2F-1,A,2F" in text and "0912345678" in text

    token = path.rsplit("/", 1)[1]
    assert client.get("/projects", headers={"Authorization": f"Bearer {token}"}).status_code == 401  # 不能拿來登入
    assert client.get("/downloads/not-a-token").status_code == 410
    assert client.post(f"/projects/{project_id}/units/export-link").status_code == 401  # 產生連結要登入


def test_download_link_not_on_public_server(client, project_id):
    from fastapi.testclient import TestClient

    from app.public_main import app as public_app

    link = client.post(f"/projects/{project_id}/units/export-link", headers=ADMIN).json()
    with TestClient(public_app) as pc:
        assert pc.get("/" + link["url"].split("://", 1)[1].split("/", 1)[1]).status_code == 404


@pytest.mark.parametrize("content", [
    "戶別*,屋主姓名,屋主電話,交屋日期\r\nD-1,張三,0911222333,2026/12/01\r\n".encode("utf-8-sig"),
    "戶別,屋主姓名,屋主電話,交屋日期\nD-1,張三,0911222333,2026-12-01\n".encode("big5"),
    "戶別\t屋主姓名\t屋主電話\t交屋日期\nD-1\t張三\t0911222333\t2026-12-01\n".encode("utf-8"),
])
def test_import_csv_variants(client, project_id, content):
    r = upload(client, project_id, content).json()
    assert r["errors"] == [] and r["created"] == ["D-1"]
    d1 = next(u for u in client.get(f"/projects/{project_id}/units", headers=ADMIN).json() if u["unit_no"] == "D-1")
    assert (d1["owner_name"], d1["owner_phone"], d1["handover_date"]) == ("張三", "0911222333", "2026-12-01")


def test_old_xls_rejected(client, project_id):
    r = upload(client, project_id, b"\xd0\xcf\x11\xe0" + b"\x00" * 100)
    assert r.status_code == 422 and ".xls" in r.json()["detail"]
